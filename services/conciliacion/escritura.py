"""Escritura de resultados de la conciliación en el libro de Excel."""
import re

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.config import DUPLICADO_FILL_COLOR, OUTPUT_SHEETS_TO_HIDE, RED_FILL_COLOR
from core.errors import ErrorSistema, ErrorUsuario
from shared.excel_io import escritor_excel
from shared.texto import normalizar_nit

COLUMNAS_INYECTADAS = [
    ('AUTORRETENCION', 'auto', "cc99ff", None),
    ('BASE', 'base', "339966", '#,##0.00'),
    ('BASE_2', 'base_2', "33EAEA", '#,##0.00'),
    ('IVA', 'iva', "FF9900", '#,##0.00'),
    ('OTROS IMPUESTOS', 'otros', "FFC000", '#,##0.00'),
    ('RETENCIONES', 'ret', "F4B183", '#,##0.00'),
]

COLOR_MOTIVOS = {
    'DIAN - Gasto Personal': 'FFCCFF',
    'DIAN - Application Response (Evento)': 'D9D9D9',
    'DIAN - Sin Num.Ext': 'F2DCDB',
    'CONT - Sin Num.Ext': 'F2DCDB',
    'DIAN - Sin Pareja en Contabilidad': 'FFF2CC',
    'CONT - Sin Pareja en DIAN': 'E2EFDA',
    'DIAN - DUPLICADO': 'FF6B6B',
    'CONT - DUPLICADO': 'FF6B6B',
    'DIAN - Sin Pareja en Contabilidad ni Devoluciones': 'FFF2CC',
    'DEVOL - Sin Pareja en DIAN': 'E2EFDA',
    'DEVOL - Sin Num.Ext': 'F2DCDB',
}

COLORES_PREFIJO = {  # (fila con fecha, fila sin fecha)
    'FC': ('B8E3BD', 'DAEDDC'),
    'GC': ('ADB3DB', 'D5D8F0'),
    'DS': ('E6BE9A', 'F2D7BF'),
    'DC': ('cddbb4', 'e7f5ce'),
    'CG': ('f5d6e7', 'f9e6f2'),
}
COLOR_DEFECTO = ("b5c0e8", "d3daf2")


def _fill(color):
    return PatternFill(start_color=color, end_color=color, fill_type="solid")


THIN_BORDER = Border(*(Side(style='thin', color='D9D9D9'),) * 4)
FILL_ROJO = _fill(RED_FILL_COLOR)
FILL_AMARILLO = _fill('FADC46')
FONT_ROJA = Font(name="Arial", size=10, color="FF0000")
FONT_NORMAL = Font(name="Arial", size=10)
FONT_NEGRITA = Font(bold=True)


def _valor(lista, i, defecto=False):
    return lista[i] if i < len(lista) else defecto


def _indice(df, columna):
    return df.columns.get_loc(columna) + 1 if columna in df.columns else None


def _diccionario_autorretenedores(df_auto):
    cols = df_auto.columns
    col_nit = next((c for c in cols if 'NIT' in str(c).upper()), None)
    col_com = next((c for c in cols if 'COMENT' in str(c).upper()), None) or \
        next((c for c in cols if 'NOMBRE' in str(c).upper()), None)
    if not (col_nit and col_com):
        return {}
    dic = {}
    for nit, comentario in zip(df_auto[col_nit], df_auto[col_com]):
        nit = normalizar_nit(nit)
        if nit:
            dic[nit] = str(comentario).strip() or 'AUTORRETENEDOR'
    return dic


def _clave_orden_auditoria(row):
    """Autorretenedores → régimen simple/comentados → resto; FC, otros, DC (agrupado por tercero); base desc."""
    auto = row['auto'].upper().strip()
    if 'AUTORRETENEDOR' in auto:
        prioridad_auto = 1
    elif auto or 'REGIMEN SIMPLE' in row['tipo']:
        prioridad_auto = 2
    else:
        prioridad_auto = 3
    try:
        base = float(row['base'])
    except (ValueError, TypeError):
        base = 0.0
    prioridad_doc = 1 if row['tipo'].startswith('FC') else 3 if row['tipo'].startswith('DC') else 2
    agrupacion = row['tercero'] if prioridad_doc == 3 else ""
    return prioridad_auto, prioridad_doc, agrupacion, -abs(base), row['tercero']


def inyectar_columnas(wb, nombre_hoja, df_res, dict_auto):
    """Inserta AUTORRETENCION/BASE/BASE_2/IVA/OTROS/RETENCIONES después de DETALLE y reordena las filas de la hoja."""
    if not nombre_hoja or nombre_hoja not in wb.sheetnames or df_res is None or df_res.empty:
        return
    ws = wb[nombre_hoja]

    header_row = next(
        (r for r in range(1, 20) for c in range(1, ws.max_column + 1)
         if 'DETALLE' in str(ws.cell(row=r, column=c).value).strip().upper()),
        0,
    )
    if not header_row:
        return

    encabezado = lambda c: str(ws.cell(row=header_row, column=c).value).strip().upper()

    # Limpia columnas inyectadas en una ejecución previa
    col_detalle = next((c for c in range(1, ws.max_column + 1) if encabezado(c) == 'DETALLE'), 0)
    if col_detalle:
        nombres = {n for n, *_ in COLUMNAS_INYECTADAS}
        previas = 0
        for i in range(1, len(COLUMNAS_INYECTADAS) + 1):
            if encabezado(col_detalle + i) not in nombres:
                break
            previas += 1
        if previas:
            ws.delete_cols(col_detalle + 1, amount=previas)

    col_tercero = col_nit = col_tipo = 0
    for c in range(1, ws.max_column + 1):
        val = encabezado(c)
        if 'DETALLE' in val:
            col_tercero = c
        elif 'NIT' in val or 'C.C' in val:
            col_nit = c
        elif 'TIPO' in val:
            col_tipo = c
    if not col_tercero:
        return

    registros = df_res.to_dict('records')
    vacio = {'BASE': 0, 'BASE_2': 0, 'IVA': 0, 'OTROS IMPUESTOS': 0, 'RETENCIONES': 0}
    max_col = ws.max_column
    filas = []
    for idx, r in enumerate(range(header_row + 1, ws.max_row + 1)):
        valores = [ws.cell(row=r, column=c).value for c in range(1, max_col + 1)]
        calc = registros[idx] if idx < len(registros) else vacio
        nit = re.sub(r'\D', '', str(valores[col_nit - 1])) if col_nit else ''
        filas.append({
            'original_values': valores,
            'base': calc.get('BASE', 0),
            'base_2': calc.get('BASE_2', 0),
            'iva': calc.get('IVA', 0),
            'otros': calc.get('OTROS IMPUESTOS', 0),
            'ret': calc.get('RETENCIONES', 0),
            'auto': dict_auto.get(nit, ''),
            'tipo': str(valores[col_tipo - 1]).strip().upper() if col_tipo else '',
            'tercero': str(valores[col_tercero - 1]).strip().upper(),
        })
    filas.sort(key=_clave_orden_auditoria)

    idx_insert = col_tercero + 1
    n_iny = len(COLUMNAS_INYECTADAS)
    ws.insert_cols(idx_insert, amount=n_iny)
    for offset, (titulo, _, color, _) in enumerate(COLUMNAS_INYECTADAS):
        celda = ws.cell(row=header_row, column=idx_insert + offset)
        celda.value = titulo
        celda.fill = _fill(color)

    for i, fila in enumerate(filas):
        r = header_row + 1 + i
        orig = fila['original_values']
        for c in range(1, idx_insert):
            ws.cell(row=r, column=c).value = orig[c - 1]
        for offset, (_, clave, _, formato) in enumerate(COLUMNAS_INYECTADAS):
            celda = ws.cell(row=r, column=idx_insert + offset)
            celda.value = fila[clave]
            if formato:
                celda.number_format = formato
        for c in range(idx_insert, len(orig) + 1):
            ws.cell(row=r, column=c + n_iny).value = orig[c - 1]


def _formato_numerico(hojas_df):
    for ws, df in hojas_df:
        for col in ('BASE', 'BASE_2', 'IVA', 'OTROS IMPUESTOS', 'RETENCIONES', 'Total'):
            if col in df.columns:
                c = df.columns.get_loc(col) + 1
                for r in range(2, len(df) + 2):
                    ws.cell(row=r, column=c).number_format = '#,##0.00'


def _color_fila(motivo, incompleta, prefijo, tiene_fecha):
    llenado, vacio = COLORES_PREFIJO.get(prefijo, COLOR_DEFECTO) if not incompleta else COLOR_DEFECTO
    if motivo:
        if 'DUPLICADO' in motivo:
            return DUPLICADO_FILL_COLOR
        if 'PERSONAL' in motivo.upper():
            return 'FFCCFF'
        if 'No coincide tercero' in motivo:
            return llenado
        return COLOR_MOTIVOS.get(motivo, 'FFC7CE')
    return llenado if tiene_fecha else vacio


def _formatear_dian_vs_cont(ws, df, res):
    max_row, max_col = len(df) + 1, len(df.columns)
    ws.freeze_panes = 'A2'

    header_fill = _fill("366092")
    header_font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    for c in range(1, max_col + 1):
        cell = ws.cell(row=1, column=c)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = THIN_BORDER

    ci = {k: _indice(df, k) for k in ('Fecha', 'Novedad', 'Tipo de documento', 'BASE', 'IVA', 'Num.Ext', 'NIT Emisor')}
    fills = {}

    for r in range(2, max_row + 1):
        i = r - 2
        fecha = ws.cell(row=r, column=ci['Fecha']).value if ci['Fecha'] else None
        motivo = ws.cell(row=r, column=ci['Novedad']).value if ci['Novedad'] else ''
        motivo = str(motivo).strip() if motivo and str(motivo).strip() else ''
        color = _color_fila(
            motivo, _valor(res.incompletas, i, True), _valor(res.prefix_doc_group, i, ''),
            fecha is not None and str(fecha).strip() != '',
        )
        fill = fills.setdefault(color, _fill(color))

        tipo_doc = ws.cell(row=r, column=ci['Tipo de documento']).value if ci['Tipo de documento'] else ''
        font = FONT_ROJA if 'nota de crédito electrónica' in str(tipo_doc).strip().lower() else FONT_NORMAL

        for c in range(1, max_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = fill
            cell.font = font
            cell.border = THIN_BORDER

        if _valor(res.dif_base, i) and ci['BASE']:
            ws.cell(row=r, column=ci['BASE']).fill = FILL_ROJO
        if _valor(res.dif_iva, i) and ci['IVA']:
            ws.cell(row=r, column=ci['IVA']).fill = FILL_ROJO
        if _valor(res.box_color, i, '') == 'YELLOW':
            for k in ('BASE', 'IVA'):
                if ci[k]:
                    ws.cell(row=r, column=ci[k]).fill = FILL_AMARILLO
        if ci['Num.Ext'] and _valor(res.caracter_especial, i):
            ws.cell(row=r, column=ci['Num.Ext']).fill = FILL_ROJO
        if ci['NIT Emisor'] and _valor(res.dif_nit, i):
            ws.cell(row=r, column=ci['NIT Emisor']).fill = FILL_ROJO

    if ci['BASE'] and ci['IVA']:
        _recuadrar_parejas(ws, res.box_color, min(ci['BASE'], ci['IVA']), max(ci['BASE'], ci['IVA']))


def _recuadrar_parejas(ws, box_color, col_ini, col_fin):
    """Recuadro grueso sobre BASE/IVA en parejas consecutivas que cuadran (efecto cero)."""
    lado = Side(style='medium', color='000000')
    i = 0
    while i < len(box_color):
        if box_color[i] == 'BLACK' and i + 1 < len(box_color) and box_color[i + 1] == 'BLACK':
            for c in range(col_ini, col_fin + 1):
                izq = lado if c == col_ini else THIN_BORDER.left
                der = lado if c == col_fin else THIN_BORDER.right
                ws.cell(row=i + 2, column=c).border = Border(top=lado, bottom=THIN_BORDER.bottom, left=izq, right=der)
                ws.cell(row=i + 3, column=c).border = Border(top=THIN_BORDER.top, bottom=lado, left=izq, right=der)
            i += 2
        else:
            i += 1


def _escribir_totales(ws, fila, totales, ci):
    if ci['BASE'] and ci['BASE'] > 1:
        lbl = ws.cell(row=fila, column=ci['BASE'] - 1)
        lbl.value = "TOTAL CONCILIADOS:"
        lbl.font = FONT_NEGRITA
        lbl.alignment = Alignment(horizontal='right')
    for col, valor in totales.items():
        if ci[col]:
            c = ws.cell(row=fila, column=ci[col])
            c.value = valor
            c.number_format = '#,##0.00'
            c.font = FONT_NEGRITA


def _insertar_separadores_y_totales(ws, df, res):
    """Separa grupos de conciliados (emisor/prefijo) y agrega el bloque de totales antes de los no conciliados."""
    inc = res.incompletas
    if not inc:
        return
    ci = {k: _indice(df, k) for k in ('BASE', 'IVA', 'BASE_2')}
    conciliados = ~np.array(inc)
    totales = {
        col: pd.to_numeric(df.loc[conciliados, col], errors='coerce').fillna(0).sum() if col in df.columns else 0
        for col in ('BASE', 'IVA', 'BASE_2')
    }

    frontera = False
    for r in range(len(df) + 1, 2, -1):
        cur, prev = r - 2, r - 3
        if inc[cur] and not inc[prev]:
            ws.insert_rows(r, amount=5)
            frontera = True
            _escribir_totales(ws, r + 2, totales, ci)
        elif not inc[cur] and not inc[prev]:
            if (res.nombre_emisor_group[cur] != res.nombre_emisor_group[prev]
                    or res.prefix_doc_group[cur] != res.prefix_doc_group[prev]):
                ws.insert_rows(r, amount=1)

    if not frontera and not all(inc):
        fin = ws.max_row + 1
        ws.insert_rows(fin, amount=5)
        _escribir_totales(ws, fin + 2, totales, ci)


def _ajustar_anchos(ws, df):
    for idx, col_name in enumerate(df.columns, start=1):
        letra = get_column_letter(idx)
        nombre = str(col_name).strip().upper()
        if nombre == "CUFE/CUDE":
            ws.column_dimensions[letra].width = 6
        elif nombre == "NOMBRE EMISOR":
            ws.column_dimensions[letra].width = 26
        else:
            largo = df[col_name].astype(str).str.len().max()
            largo = 0 if pd.isna(largo) else int(largo)
            ws.column_dimensions[letra].width = max(largo, len(str(col_name)), 12) + 3


def escribir_resultados(file_path, sheet_names, nombre_dc, df_resultado, df_auditoria, df_res_auditoria,
                        df_autorretenedores, res):
    try:
        with escritor_excel(file_path) as writer:
            df = res.df_dian_vs_cont
            df_resultado.to_excel(writer, index=False, sheet_name='Resultados')
            df_auditoria.to_excel(writer, index=False, sheet_name='auditoria')
            df.to_excel(writer, index=False, sheet_name='DIAN VS CONT')

            wb = writer.book
            ws = writer.sheets['DIAN VS CONT']

            dict_auto = _diccionario_autorretenedores(df_autorretenedores)
            inyectar_columnas(wb, sheet_names['aud_comp'], df_res_auditoria, dict_auto)
            inyectar_columnas(wb, nombre_dc, res.df_res_dc, dict_auto)

            _formato_numerico([
                (writer.sheets['Resultados'], df_resultado),
                (writer.sheets['auditoria'], df_auditoria),
                (ws, df),
            ])
            _formatear_dian_vs_cont(ws, df, res)
            _insertar_separadores_y_totales(ws, df, res)
            _ajustar_anchos(ws, df)

            for nombre in OUTPUT_SHEETS_TO_HIDE:
                if nombre in wb.sheetnames:
                    wb[nombre].sheet_state = 'hidden'
            if not any(h.sheet_state == 'visible' for h in wb.worksheets):
                wb.worksheets[0].sheet_state = 'visible'
    except (ErrorUsuario, ErrorSistema):
        raise
    except Exception as e:
        raise ErrorSistema(f"Error escribiendo resultados en el archivo Excel: {e}") from e
