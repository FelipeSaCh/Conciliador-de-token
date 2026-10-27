"""Informe de IVA: cálculo (generado vs descontable) y escritura en la hoja 'Informe IVA'."""
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.errors import ErrorSistema, ErrorUsuario, Reportador, logger
from shared.dataframes import a_numero, buscar_columna
from shared.texto import MESES_ES, nombre_archivo_seguro, normalizar_ascii

TIPOS_DOC_VALIDOS = {"Factura electrónica", "Nota credito", "Nota debito"}
GRUPOS_VALIDOS = {"Emitido", "Recibido"}
NOMBRES_CONCEPTOS = {"Emitido": "IVA GENERADO", "Recibido": "IVA DESCONTABLE"}

COLUMNA_GRUPO = "Grupo"
COLUMNA_TIPO_DOC = "Tipo de documento"
COLUMNA_IVA = "IVA"
HOJA_INFORME_IVA = "Informe IVA"

TIPOS_DOC_POR_DEFECTO_NORM = {normalizar_ascii(t) for t in TIPOS_DOC_VALIDOS}


def es_nota_credito(tipo_doc):
    t = normalizar_ascii(tipo_doc)
    return "credito" in t or "ajuste" in t


def etiqueta_saldo(valor):
    if valor > 0:
        return "VALOR A PAGAR"
    return "SALDO A FAVOR" if valor < 0 else "SIN SALDO"


class GeneradorInformeIVA(Reportador):
    def __init__(self, ruta_archivo, hoja_token, progress_callback=None, tipos_documento=None):
        self.ruta_archivo = ruta_archivo
        self.hoja_token = hoja_token
        self.progress_callback = progress_callback
        self.tipos_documento = set(tipos_documento) if tipos_documento else None
        self.df_filtrado = None
        self.resumen = None
        self.periodo_texto = ""

    @staticmethod
    def _cargar_hoja(ruta, hoja):
        try:
            return pd.read_excel(ruta, sheet_name=hoja)
        except Exception as e:
            logger.error(f"No se pudo leer la hoja '{hoja}': {e}")
            raise ErrorSistema(f"No se pudo leer la hoja '{hoja}' del archivo: {e}")

    @classmethod
    def obtener_tipos_documento(cls, ruta_archivo, hoja_token):
        df = cls._cargar_hoja(ruta_archivo, hoja_token)
        if COLUMNA_TIPO_DOC not in df.columns:
            raise ErrorUsuario(f"La hoja '{hoja_token}' no contiene la columna requerida '{COLUMNA_TIPO_DOC}'.")
        tipos = sorted({
            str(t).strip() for t in df[COLUMNA_TIPO_DOC].dropna()
            if str(t).strip() and str(t).strip().lower() != "nan"
        })
        if not tipos:
            raise ErrorUsuario("No se encontraron tipos de documento en la hoja seleccionada.")
        return tipos, [t for t in tipos if normalizar_ascii(t) in TIPOS_DOC_POR_DEFECTO_NORM]

    def generar(self):
        self._reportar(f"Leyendo hoja de token: {self.hoja_token}")
        df = self._cargar_hoja(self.ruta_archivo, self.hoja_token)

        for col in (COLUMNA_GRUPO, COLUMNA_TIPO_DOC, COLUMNA_IVA):
            if col not in df.columns:
                raise ErrorUsuario(f"La hoja '{self.hoja_token}' no contiene la columna requerida '{col}'.")

        df[COLUMNA_GRUPO] = df[COLUMNA_GRUPO].astype(str).str.strip()
        df[COLUMNA_TIPO_DOC] = df[COLUMNA_TIPO_DOC].astype(str).str.strip()
        df[COLUMNA_IVA] = a_numero(df[COLUMNA_IVA])

        tipos_incluidos = self.tipos_documento or TIPOS_DOC_VALIDOS
        self._reportar(f"Tipos de documento incluidos: {', '.join(sorted(tipos_incluidos))}")

        filtro = df[COLUMNA_TIPO_DOC].isin(tipos_incluidos) & df[COLUMNA_GRUPO].isin(GRUPOS_VALIDOS)

        # IVA de gastos personales no es descontable: se excluye solo en Recibidos
        if "CONCEPTO" in df.columns:
            excluir = (df[COLUMNA_GRUPO] == "Recibido") & df["CONCEPTO"].astype(str).str.strip().str.upper().str.contains("PERSONAL", na=False)
            filtro &= ~excluir
            if excluir.sum() > 0:
                self._reportar(f"Se ignoraron {excluir.sum()} registro(s) de 'Recibido' por tener conceptos personales.")

        df_filtrado = df[filtro].copy()
        if df_filtrado.empty:
            raise ErrorUsuario(
                "No se encontraron registros que cumplan con los criterios indicados, "
                "o todos fueron filtrados/excluidos."
            )
        self._reportar(f"Registros filtrados: {len(df_filtrado)}")

        df_filtrado["Signo"] = self._calcular_signo(df_filtrado)
        df_filtrado["IVA Ajustado"] = df_filtrado[COLUMNA_IVA] * df_filtrado["Signo"]

        self.df_filtrado = df_filtrado
        self.resumen = self._construir_resumen(df_filtrado)
        self.periodo_texto = self._obtener_texto_periodo(df)

        self._reportar("Escribiendo informe en el archivo de Excel...")
        self._escribir_informe(df_filtrado, self.resumen, self.periodo_texto)
        self._reportar("Informe de IVA generado correctamente.")
        return self.resumen

    @staticmethod
    def _calcular_signo(df):
        """Emitido: factura +, nota crédito −. Recibido: factura −, nota crédito +."""
        nota = df[COLUMNA_TIPO_DOC].map(es_nota_credito)
        grupo = df[COLUMNA_GRUPO]
        return np.select(
            [grupo == "Emitido", grupo == "Recibido"],
            [np.where(nota, -1, 1), np.where(nota, 1, -1)],
            default=0,
        )

    @staticmethod
    def _obtener_texto_periodo(df):
        col_fecha = buscar_columna(df.columns, "fecha emisión", "fecha emision", "fecha_emision")
        if not col_fecha:
            return ""
        fechas = pd.to_datetime(df[col_fecha], errors="coerce", dayfirst=True).dropna()
        if fechas.empty:
            return ""
        meses = [MESES_ES[m] for m in sorted(fechas.dt.month.unique()) if m in MESES_ES]
        if not meses:
            return ""
        anios = sorted(fechas.dt.year.unique())
        texto_meses = meses[0] if len(meses) == 1 else f"{meses[0]} - {meses[-1]}"
        texto_anios = str(anios[0]) if len(anios) == 1 else f"{anios[0]}-{anios[-1]}"
        return f"PERÍODO: {texto_meses.upper()} {texto_anios}"

    def _construir_resumen(self, df):
        resumen = {"grupos": {}, "totales": {}}
        for grupo in GRUPOS_VALIDOS:
            df_grupo = df[df[COLUMNA_GRUPO] == grupo]
            notas = df_grupo[COLUMNA_TIPO_DOC].map(es_nota_credito).astype(bool)
            resumen["grupos"][grupo] = {
                "facturas_cantidad": int((~notas).sum()),
                "facturas_iva": round(df_grupo.loc[~notas, COLUMNA_IVA].sum(), 2),
                "notas_cantidad": int(notas.sum()),
                "notas_iva": round(df_grupo.loc[notas, COLUMNA_IVA].sum(), 2),
                "iva_neto": round(df_grupo["IVA Ajustado"].sum(), 2),
            }
        iva_emitido = resumen["grupos"]["Emitido"]["iva_neto"]
        iva_recibido = resumen["grupos"]["Recibido"]["iva_neto"]
        resumen["totales"] = {
            "iva_emitido": iva_emitido,
            "iva_recibido": iva_recibido,
            "iva_a_pagar": round(iva_emitido + iva_recibido, 2),
            "total_registros": len(df),
        }
        return resumen

    def nombre_sugerido_pdf(self):
        cliente = ""
        df = self.df_filtrado
        if df is not None and {"Grupo", "Nombre Emisor"} <= set(df.columns):
            emitidos = df[df["Grupo"].astype(str).str.strip() == "Emitido"]
            if not emitidos.empty:
                val = emitidos["Nombre Emisor"].iloc[0]
                if pd.notna(val) and str(val).strip():
                    cliente = str(val).strip() + " - "
        periodo = self.periodo_texto.replace("PERÍODO:", "").strip() if self.periodo_texto else ""
        sufijo = f" - {periodo}" if periodo else ""
        return nombre_archivo_seguro(f"{cliente}INFORME DE IVA{sufijo}.pdf")

    def generar_pdf(self, ruta_pdf=None):
        if self.df_filtrado is None or self.resumen is None:
            raise ErrorSistema("Debes generar el informe de IVA antes de exportarlo a PDF.")
        from services.reportes.informe_iva_pdf import construir_pdf_iva

        ruta_pdf = ruta_pdf or str(Path(self.ruta_archivo).with_name(self.nombre_sugerido_pdf()))
        self._reportar("Generando documento PDF...")
        construir_pdf_iva(ruta_pdf, self.df_filtrado, self.resumen, self.periodo_texto, self.hoja_token)
        self._reportar("PDF generado correctamente.")
        return ruta_pdf

    def _escribir_informe(self, df, resumen, periodo_texto=""):
        try:
            libro = load_workbook(self.ruta_archivo)
        except Exception as e:
            raise ErrorSistema(f"No se pudo abrir el archivo para escribir el informe: {e}")

        if HOJA_INFORME_IVA in libro.sheetnames:
            del libro[HOJA_INFORME_IVA]
        hoja = libro.create_sheet(HOJA_INFORME_IVA, 0)

        font_seccion = Font(size=11, bold=True, color="FFFFFF")
        fill_seccion = PatternFill("solid", fgColor="2563EB")
        font_enc = Font(bold=True, color="FFFFFF")
        fill_enc = PatternFill("solid", fgColor="374151")
        fill_total = PatternFill("solid", fgColor="D1FAE5")
        font_total = Font(bold=True)
        borde = Border(*(Side(style="thin", color="D3D8E0"),) * 4)
        moneda = "#,##0.00"
        fill_grupo = {"Emitido": PatternFill("solid", fgColor="EDFFC3"), "Recibido": PatternFill("solid", fgColor="E0E7FF")}

        def seccion(fila, texto, n_cols):
            hoja.cell(row=fila, column=1, value=texto).font = font_seccion
            hoja.merge_cells(start_row=fila, start_column=1, end_row=fila, end_column=n_cols)
            for c in range(1, n_cols + 1):
                hoja.cell(row=fila, column=c).fill = fill_seccion

        def encabezados(fila, textos):
            for c, texto in enumerate(textos, start=1):
                celda = hoja.cell(row=fila, column=c, value=texto)
                celda.font, celda.fill, celda.border = font_enc, fill_enc, borde

        hoja.cell(row=1, column=1, value="INFORME DE IVA").font = Font(size=14, bold=True, color="1F2937")
        if periodo_texto:
            hoja.cell(row=1, column=3, value=periodo_texto).font = Font(size=11, bold=True, color="2563EB")
        hoja.cell(row=2, column=1, value=f"Hoja origen: {self.hoja_token}").font = Font(italic=True, color="5B6472")

        fila = 4
        seccion(fila, "RESUMEN GENERAL", 4)
        fila += 1
        encabezados(fila, ["Concepto", "Facturas y ND (cant. / IVA)", "Notas Crédito (cant. / IVA)", "IVA Neto"])
        fila += 1

        for grupo in ("Emitido", "Recibido"):
            d = resumen["grupos"][grupo]
            valores = [
                NOMBRES_CONCEPTOS[grupo],
                f"{d['facturas_cantidad']} / {d['facturas_iva']:,.2f}",
                f"{d['notas_cantidad']} / {d['notas_iva']:,.2f}",
                d["iva_neto"],
            ]
            for c, v in enumerate(valores, start=1):
                hoja.cell(row=fila, column=c, value=v).border = borde
            hoja.cell(row=fila, column=4).number_format = moneda
            fila += 1

        total = resumen["totales"]["iva_a_pagar"]
        hoja.cell(row=fila, column=1, value=etiqueta_saldo(total))
        hoja.cell(row=fila, column=4, value=total).number_format = moneda
        for c in (1, 4):
            hoja.cell(row=fila, column=c).font = font_total
            hoja.cell(row=fila, column=c).fill = fill_total
        for c in range(1, 5):
            hoja.cell(row=fila, column=c).border = borde
        fila += 3

        columnas = list(df.columns)
        seccion(fila, "DETALLE DE REGISTROS", len(columnas))
        fila += 1
        encabezados(fila, columnas)
        fila_inicio_detalle = fila
        fila += 1

        cols_moneda = {i for i, c in enumerate(columnas, start=1) if c in (COLUMNA_IVA, "IVA Ajustado")}
        idx_grupo = columnas.index(COLUMNA_GRUPO)
        for registro in df.itertuples(index=False, name=None):
            relleno = fill_grupo.get(registro[idx_grupo])
            for c, valor in enumerate(registro, start=1):
                celda = hoja.cell(row=fila, column=c, value=valor)
                celda.border = borde
                if relleno:
                    celda.fill = relleno
                if c in cols_moneda:
                    celda.number_format = moneda
            fila += 1

        minimos = {1: 28, 2: 24, 3: 24, 4: 18}
        for c, col in enumerate(columnas, start=1):
            max_len = max([len(str(col))] + df[col].astype(str).str.len().tolist())
            ancho = min(max(max_len + 2, 12), 45)
            hoja.column_dimensions[get_column_letter(c)].width = max(ancho, minimos.get(c, 0))

        hoja.freeze_panes = hoja.cell(row=fila_inicio_detalle + 1, column=1)
        try:
            libro.save(self.ruta_archivo)
        except Exception as e:
            raise ErrorSistema(f"No se pudo guardar el informe en el archivo: {e}")
