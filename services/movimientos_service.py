"""Reestructuración (pivoteo) de hojas de movimientos contables hacia la hoja MOVS_AUD."""
from pathlib import Path
from typing import Callable, Dict, List, Optional, Tuple, Union

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, NamedStyle, PatternFill, Side
from openpyxl.utils import get_column_letter

from core.errors import ErrorSistema, ErrorUsuario, Reportador, logger
from shared.cuentas import ordenar_por_prefijos
from shared.excel_io import escritor_excel, leer_excel
from shared.texto import clave_columna, limpiar_codigo

MESES_CORTOS = {i: m for i, m in enumerate(
    ['Ene', 'Feb', 'Mar', 'Abr', 'May', 'Jun', 'Jul', 'Ago', 'Sep', 'Oct', 'Nov', 'Dic'], start=1)}
MESES_CORTOS.update({str(k): v for k, v in list(MESES_CORTOS.items())})
MESES_CORTOS.update({f"{k:02d}": v for k, v in list(MESES_CORTOS.items()) if isinstance(k, int) and k < 10})

RENOMBRES = {
    "TIPO": "Tipo", "NUMERO": "Número", "NO.EXT": "Num.Ext", "Nit/C.C.": "NIT", "NIT": "Nit/C.C.",
    "Nombre Emisor": "Tercero", "TERCERO": "Tercero", "AÑO": "Año", "MES": "Mes", "DIA": "Dia",
}
COLUMNAS_DESCARTADAS = ["CHEQ.NO", "ELABORO", "SUC.PTO", "CCOSTO", "SCCOSTO", "REF1", "REF2", "REF3", "REF4",
                        "No. AUTORIZACION(DATAFONO)"]
ORDEN_INDICE = ["Tipo", "Número", "Num.Ext", "Fecha", "Tercero", "DETALLE", "DOCRELA", "NIT"]
PRIORIDAD_CUENTAS = ('2365', '2367', '2368', '15', '6', '5', '22')
PRIORIDAD_DETALLE = ('6', '5', '7', '15')
TIPOS_AL_FINAL = ('GC', 'GS', 'CG')


class TransformadorMovimientos(Reportador):
    PREFIJOS_PERMITIDOS = ("FC", "GS", "DC", "DS", "GC", "CG", "GI")
    HOJA_DESTINO = "MOVS_AUD"
    COLUMNA_FECHA_ORIGEN = "__FECHA_ORIGEN__"

    # Formatos detectados por columnas requeridas (claves sin tildes/espacios/signos) → nombres estándar.
    FORMATOS: Dict[str, dict] = {
        "estandar": {
            "descripcion": "Movimientos estándar",
            "requeridas": {
                "DEBITO": ("debito",), "CREDITO": ("credito",), "CUENTA": ("cuenta",),
                "NOM. CUENTA": ("nomcuenta",), "TIPO": ("tipo",),
            },
            "opcionales": {},
            "detalle_vacio": False,
            "eliminar": (),
        },
        "auxiliar_documento": {
            "descripcion": "Auxiliar por documento",
            "requeridas": {
                "DEBITO": ("debitos", "dbitos"), "CREDITO": ("creditos", "crditos"), "CUENTA": ("auxiliar",),
                "NOM. CUENTA": ("descauxiliar",), "TIPO": ("tipodocto",), "NO.EXT": ("docto",),
            },
            "opcionales": {
                "TERCERO": ("razonsocialtercerodocto", "raznsocialtercerodocto"),
                "Nit/C.C.": ("nittercero",),
                COLUMNA_FECHA_ORIGEN: ("fecha",),
            },
            "detalle_vacio": True,
            "eliminar": ("neto", "basegravable", "comovto"),
        },
    }

    def __init__(
        self,
        ruta_entrada: Union[str, Path],
        hoja_origen: str,
        prefijos_permitidos: Optional[List[str]] = None,
        progress_callback: Optional[Callable[[str], None]] = None,
        avance_callback: Optional[Callable[[Optional[float]], None]] = None,
    ):
        self.ruta_entrada = Path(ruta_entrada)
        self.hoja_origen = hoja_origen
        self.prefijos_permitidos = tuple(prefijos_permitidos) if prefijos_permitidos is not None else self.PREFIJOS_PERMITIDOS
        self.progress_callback = progress_callback
        self.avance_callback = avance_callback

    def _avance(self, valor: Optional[float]) -> None:
        """0-100; None = fase sin progreso medible (barra indeterminada)."""
        if self.avance_callback:
            self.avance_callback(valor)

    @classmethod
    def detectar_formato(cls, columnas) -> Tuple[str, Dict[str, str]]:
        claves = {}
        for c in columnas:
            claves.setdefault(clave_columna(c), c)

        def buscar(alias):
            return next((claves[a] for a in alias if a in claves), None)

        for nombre, fmt in cls.FORMATOS.items():
            mapeo = {buscar(alias): destino for destino, alias in fmt["requeridas"].items()}
            if None in mapeo:
                continue
            mapeo.update({o: d for d, alias in fmt["opcionales"].items() if (o := buscar(alias)) is not None})
            return nombre, mapeo

        detalle = "; ".join(f"{f['descripcion']}: {', '.join(f['requeridas'])}" for f in cls.FORMATOS.values())
        raise ErrorUsuario(f"La hoja no coincide con ningún formato conocido. Formatos soportados -> {detalle}")

    @classmethod
    def obtener_tipos(cls, ruta, hoja) -> Tuple[List[str], str]:
        nombre, mapeo = cls.detectar_formato(leer_excel(ruta, rapido=True, sheet_name=hoja, nrows=0).columns)
        col_tipo = next(o for o, d in mapeo.items() if d == "TIPO")
        df = leer_excel(ruta, rapido=True, sheet_name=hoja, usecols=[col_tipo])
        tipos = df[col_tipo].dropna().astype(str).str.strip().str.upper().unique()
        return sorted(t for t in tipos if t), cls.FORMATOS[nombre]["descripcion"]

    def ejecutar(self) -> Path:
        try:
            self._avance(None)
            df = self._cargar_y_validar()
            self._avance(20)
            df_final, cols_indice = self._transformar_y_pivotar(df)
            self._avance(40)
            self._guardar_con_formato(df_final, cols_indice)
            self._avance(100)
            self._reportar(f"Guardado finalizado exitosamente en la hoja '{self.HOJA_DESTINO}'.")
            return self.ruta_entrada
        except ErrorUsuario:
            raise
        except Exception as e:
            logger.exception("Error durante la reestructuración de movimientos.")
            raise ErrorSistema(f"Error inesperado al procesar la hoja de movimientos: {e}") from e

    def _cargar_y_validar(self) -> pd.DataFrame:
        self._reportar("Leyendo la hoja de movimientos...")
        if not self.ruta_entrada.exists():
            raise ErrorUsuario(f"El archivo origen no existe: {self.ruta_entrada}")
        try:
            df = leer_excel(self.ruta_entrada, rapido=True, sheet_name=self.hoja_origen)
        except ValueError as e:
            raise ErrorUsuario(f"La hoja '{self.hoja_origen}' no existe en el archivo.") from e
        except Exception as e:
            raise ErrorSistema(f"No se pudo leer el archivo de Excel: {e}") from e
        return self._normalizar_formato(df)

    def _normalizar_formato(self, df: pd.DataFrame) -> pd.DataFrame:
        nombre, mapeo = self.detectar_formato(df.columns)
        fmt = self.FORMATOS[nombre]
        self._reportar(f"Formato detectado: {fmt['descripcion']}")
        if nombre == "estandar":
            return df

        df = df.drop(columns=[c for c in df.columns if clave_columna(c) in fmt["eliminar"]]).rename(columns=mapeo)
        if self.COLUMNA_FECHA_ORIGEN in df.columns:
            fechas = pd.to_datetime(df.pop(self.COLUMNA_FECHA_ORIGEN), errors="coerce", dayfirst=True)
            df["AÑO"] = fechas.dt.year.astype("Int64")
            df["MES"] = fechas.dt.month.astype("Int64")
            df["DIA"] = fechas.dt.day.astype("Int64")
        if fmt["detalle_vacio"] and "DETALLE" not in df.columns:
            df["DETALLE"] = ""
        return df

    def _filtrar(self, df):
        if "Tipo" not in df.columns:
            raise ErrorUsuario("El archivo no contiene la columna 'TIPO' o 'Tipo'.")
        if not self.prefijos_permitidos:
            raise ErrorUsuario("No se seleccionó ningún tipo de documento permitido.")

        df = df[df["Tipo"].astype(str).str.strip().str.upper().str.startswith(self.prefijos_permitidos)]
        if df.empty:
            raise ErrorUsuario(
                f"No se encontraron registros que inicien con los tipos seleccionados: {', '.join(self.prefijos_permitidos)}"
            )
        if "DETALLE" in df.columns:
            df = df[~df["DETALLE"].astype(str).str.contains("ANULADO", case=False, na=False)]
        return df.drop(columns=[c for c in COLUMNAS_DESCARTADAS if c in df.columns]).copy()

    @staticmethod
    def _consolidar_fecha(df):
        if all(c in df.columns for c in ("Año", "Mes", "Dia")):
            mes = df["Mes"].map(MESES_CORTOS).fillna(df["Mes"].astype(str))
            dia = limpiar_codigo(df["Dia"].astype(str), strip=False).str.zfill(2)
            anio = limpiar_codigo(df["Año"].astype(str), strip=False)
            df["Fecha"] = mes + "-" + dia + "-" + anio
            df.drop(columns=["Año", "Mes", "Dia"], inplace=True)

    def _transformar_y_pivotar(self, df: pd.DataFrame) -> Tuple[pd.DataFrame, List[str]]:
        self._reportar("Estandarizando nombres de columnas y limpiando basura...")
        df = self._filtrar(df.rename(columns=RENOMBRES))
        self._consolidar_fecha(df)

        if "Num.Ext" in df.columns:
            df["Num.Ext"] = df["Num.Ext"].astype(str).str.replace(r'\.0$', '', regex=True).str.strip()
            df["Num.Ext"] = df["Num.Ext"].replace(['nan', 'None', '<NA>'], '')

        self._reportar("Aplicando cálculos numéricos y unificando registros duplicados...")
        df["DEBITO"] = pd.to_numeric(df["DEBITO"], errors="coerce").fillna(0)
        df["CREDITO"] = pd.to_numeric(df["CREDITO"], errors="coerce").fillna(0)
        df["VALOR_NETO"] = df["DEBITO"] - df["CREDITO"]
        df["CUENTA_CONCAT"] = df["CUENTA"].astype(str) + " - " + df["NOM. CUENTA"].astype(str)

        claves = [c for c in ("Tipo", "Número", "Num.Ext") if c in df.columns]
        cols_sistema = {"CUENTA", "NOM. CUENTA", "CUENTA_CONCAT", "DEBITO", "CREDITO", "VALOR_NETO", "DETALLE"}
        cols_desc = [c for c in df.columns if c not in claves and c not in cols_sistema]

        df_final = (
            df.groupby(claves + ["CUENTA_CONCAT"])["VALOR_NETO"].sum()
            .unstack("CUENTA_CONCAT", fill_value=0).reset_index()
        )
        df_final.columns.name = None
        self._avance(30)

        if cols_desc:
            df_final = df_final.merge(df.groupby(claves, as_index=False)[cols_desc].first(), on=claves, how="left")
        if "DETALLE" in df.columns:
            # DETALLE del documento: prioriza la línea de cuenta de gasto/costo/activo
            det = df[claves + ["CUENTA", "DETALLE"]].copy()
            det["_p"] = 2 - det["CUENTA"].astype(str).str.strip().str.startswith(PRIORIDAD_DETALLE).astype(int)
            det = det.sort_values(by=claves + ["_p"], kind="stable").drop_duplicates(subset=claves, keep="first")
            df_final = df_final.merge(det[claves + ["DETALLE"]], on=claves, how="left")
            cols_desc.append("DETALLE")

        cols_indice = [c for c in ORDEN_INDICE if c in df_final.columns]
        cols_indice += [c for c in df_final.columns if c in cols_desc and c not in cols_indice]
        cols_cuentas = ordenar_por_prefijos([c for c in df_final.columns if c not in cols_indice], PRIORIDAD_CUENTAS)
        df_final = df_final[cols_indice + cols_cuentas]

        if "Tipo" in df_final.columns and "Tercero" in df_final.columns:
            self._reportar("Ordenando filas: GC, GS, CG al final y agrupados por Tercero...")
            al_final = df_final["Tipo"].astype(str).str.strip().str.upper().str.startswith(TIPOS_AL_FINAL)
            df_final = pd.concat(
                [df_final[~al_final], df_final[al_final].sort_values(by="Tercero")], ignore_index=True
            )

        return df_final, cols_indice

    def _guardar_con_formato(self, df_final: pd.DataFrame, cols_indice: List[str]) -> None:
        self._reportar(f"Guardando resultados en la hoja '{self.HOJA_DESTINO}' del archivo original...")

        borde_fino = Side(border_style="thin", color="E2E8F0")
        borde_enc = Side(border_style="medium", color="0F172A")
        cuerpo = Border(left=borde_fino, right=borde_fino, top=borde_fino, bottom=borde_fino)
        font_body = Font(name="Arial", size=10, color="000000")
        par, impar = (PatternFill(start_color=c, end_color=c, fill_type="solid") for c in ("FFFFFF", "F8FAFC"))
        izq = Alignment(horizontal="left", vertical="center")
        der = Alignment(horizontal="right", vertical="center")
        moneda = '"$" #,##0.00'
        estilos = {
            "movs_aud_hdr": (PatternFill(start_color="ffcc00", end_color="ffcc00", fill_type="solid"),
                             Font(name="Arial", size=11, bold=True, color="000000"),
                             Border(left=borde_fino, right=borde_fino, top=borde_enc, bottom=borde_enc),
                             Alignment(horizontal="center", vertical="center", wrap_text=True), "General"),
            "movs_aud_txt_par": (par, font_body, cuerpo, izq, "General"),
            "movs_aud_txt_impar": (impar, font_body, cuerpo, izq, "General"),
            "movs_aud_num_par": (par, font_body, cuerpo, der, moneda),
            "movs_aud_num_impar": (impar, font_body, cuerpo, der, moneda),
        }

        self._avance(None)
        with escritor_excel(self.ruta_entrada) as writer:
            df_final.to_excel(writer, index=False, sheet_name=self.HOJA_DESTINO)
            ws = writer.sheets[self.HOJA_DESTINO]
            self._avance(55)

            existentes = set(writer.book.named_styles)
            for nombre, (fill, font, border, align, num_fmt) in estilos.items():
                if nombre not in existentes:
                    writer.book.add_named_style(NamedStyle(
                        name=nombre, fill=fill, font=font, border=border, alignment=align, number_format=num_fmt))

            max_row, max_col = df_final.shape[0] + 1, df_final.shape[1]
            n_indice = len(cols_indice)
            max_lineas = 1
            for idx, col in enumerate(df_final.columns, start=1):
                letra = get_column_letter(idx)
                ws.cell(row=1, column=idx).style = "movs_aud_hdr"
                largo_titulo = len(str(col))
                if idx <= n_indice:
                    largo = df_final[col].astype(str).str.len().max() if not df_final.empty else 0
                    largo = int(largo) if pd.notna(largo) else 0
                    ws.column_dimensions[letra].width = min(max(max(largo, largo_titulo) + 3, 12), 45)
                    lineas = largo_titulo // 35 + 1
                else:
                    ws.column_dimensions[letra].width = 30
                    lineas = largo_titulo // 30 + 1
                max_lineas = max(max_lineas, lineas)
            ws.row_dimensions[1].height = min(max(24, max_lineas * 14), 32)

            total = max(max_row - 1, 1)
            paso = max(total // 100, 1)
            for r, fila in enumerate(ws.iter_rows(min_row=2, max_row=max_row, max_col=max_col), start=2):
                if (r - 2) % paso == 0:
                    self._avance(55 + 40 * (r - 1) / total)
                sufijo = "par" if r % 2 == 0 else "impar"
                ws.row_dimensions[r].height = 20
                for celda in fila[:n_indice]:
                    celda.style = f"movs_aud_txt_{sufijo}"
                for celda in fila[n_indice:]:
                    celda.style = f"movs_aud_num_{sufijo}"

            ws.freeze_panes = "A2"
            self._reportar("Escribiendo archivo en disco...")
            self._avance(None)
