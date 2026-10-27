from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table

from core.errors import ColumnaFaltanteError, ErrorSistema, ErrorUsuario, HojaNoEncontradaError, Reportador, logger
from shared.dataframes import buscar_columna, limpiar_encabezados
from shared.excel_io import escritor_excel, hojas_de, leer_hoja_por_palabra_clave
from shared.texto import limpiar_codigo, normalizar_ascii

COLUMNAS_CALCULADAS = ["CONCEPTO", "TERCERO", "TIPO", "TIPO-DETALLE", "BASE", "Num.Ext"]
COLUMNAS_OCULTAS = {'TIPO', 'TIPO-DETALLE', 'BASE', 'Num.Ext'}
COLUMNAS_SOBRESCRIBIR = {'CONCEPTO', 'TERCERO', 'TIPO', 'TIPO-DETALLE', 'NIT_CRUCE'}


class FormateadorToken(Reportador):
    COLUMNAS_REQUERIDAS_PRINCIPAL = ['NIT Emisor', 'NIT Receptor', 'Grupo', 'Prefijo', 'Folio', 'Total', 'IVA']
    COLUMNAS_REQUERIDAS_NIT = ['NIT']

    def __init__(self, file_path, sheet_names, progress_callback=None):
        self.file_path = Path(file_path)
        self.sheet_names = sheet_names
        self.progress_callback = progress_callback

    @staticmethod
    def _validar_columnas(df, columnas, nombre_hoja):
        for col in columnas:
            if col not in df.columns:
                raise ColumnaFaltanteError(nombre_hoja, col)

    def _cargar_hoja_nit(self, clave, hojas_excel):
        """Carga contabilidad/terceros: encabezado por palabra clave 'NIT'. Si falla, DataFrame vacío."""
        nombre = self.sheet_names.get(clave)
        if not nombre or nombre not in hojas_excel:
            return pd.DataFrame(columns=self.COLUMNAS_REQUERIDAS_NIT)
        try:
            df = limpiar_encabezados(leer_hoja_por_palabra_clave(self.file_path, nombre, "NIT"), mayus=True)
            self._validar_columnas(df, self.COLUMNAS_REQUERIDAS_NIT, nombre)
            df = df.dropna(subset=['NIT'])
            return df[df['NIT'].astype(str).str.strip() != '']
        except Exception as e:
            logger.warning(f"Error al procesar hoja de {clave} '{nombre}': {e}")
            return pd.DataFrame(columns=self.COLUMNAS_REQUERIDAS_NIT)

    def ejecutar(self):
        if not self.file_path.exists():
            raise ErrorUsuario(f"El archivo no existe: {self.file_path}")
        try:
            hojas_excel = hojas_de(self.file_path)
        except Exception as e:
            raise ErrorSistema(f"No se pudo abrir el archivo Excel: {e}") from e

        nombre_principal = self.sheet_names.get("principal")
        if not nombre_principal:
            raise ErrorUsuario("No se especificó el nombre de la hoja 'principal'.")
        if nombre_principal not in hojas_excel:
            raise HojaNoEncontradaError(nombre_principal, hojas_excel)

        self._reportar("Cargando hoja principal (Token)...")
        try:
            df_principal = limpiar_encabezados(pd.read_excel(self.file_path, sheet_name=nombre_principal))
            df_principal = df_principal.dropna(how="all").reset_index(drop=True)
            self._validar_columnas(df_principal, self.COLUMNAS_REQUERIDAS_PRINCIPAL, nombre_principal)
        except Exception as e:
            raise ErrorSistema(f"Error leyendo la hoja principal: {e}") from e
        columnas_originales = df_principal.columns.tolist()

        df_conta = self._cargar_hoja_nit("contabilidad", hojas_excel)
        df_tercer = self._cargar_hoja_nit("terceros", hojas_excel)

        self._reportar("Cruzando datos entre hojas...")
        df_full = self._cruzar_datos(df_principal, df_conta, df_tercer)

        self._reportar("Calculando bases, totales y consecutivos...")
        es_personales = self._calcular_campos(df_full)

        self._reportar("Ordenando registros por Grupo, Tipo de Documento y Personales...")
        df_export = df_full[self._orden_columnas(columnas_originales)].copy()
        df_export["_ES_PERSONAL"] = es_personales.values
        df_export = self._ordenar_jerarquico(df_export)
        es_personales_ordenado = df_export.pop("_ES_PERSONAL")

        self._reportar("Escribiendo hoja Token formateada...")
        self._escribir_excel(nombre_principal, df_export, es_personales_ordenado)

        self._reportar("Formateo de Token completado con éxito.")
        return {
            "filas_procesadas": len(df_export),
            "filas_personales": int(es_personales_ordenado.sum()),
        }

    @staticmethod
    def _orden_columnas(columnas_originales):
        """Las calculadas que ya existían conservan su posición; las nuevas van al final."""
        calc_upper = [c.upper() for c in COLUMNAS_CALCULADAS]
        cols, insertadas = [], set()
        for col in columnas_originales:
            up = str(col).upper()
            if up in calc_upper:
                nombre = COLUMNAS_CALCULADAS[calc_upper.index(up)]
                cols.append(nombre)
                insertadas.add(nombre)
            else:
                cols.append(col)
        return cols + [c for c in COLUMNAS_CALCULADAS if c not in insertadas]

    @staticmethod
    def _ordenar_jerarquico(df):
        """Grupo (Emitido, Recibido, otros) → personales al final de cada grupo → tipo de documento."""
        if "Grupo" in df.columns:
            grupo = df["Grupo"].fillna("").astype(str).str.strip().str.upper()
            df["_ORDEN_GRUPO"] = np.select([grupo.str.contains("EMITIDO"), grupo.str.contains("RECIBIDO")], [0, 1], default=2)
        else:
            df["_ORDEN_GRUPO"] = 0

        col_tipo_doc = buscar_columna(df.columns, "TIPO DE DOCUMENTO", "TIPO DOCUMENTO")
        sort_cols = ["_ORDEN_GRUPO", "_ES_PERSONAL", "_ORDEN_DOC"]
        if col_tipo_doc:
            doc = df[col_tipo_doc].fillna("").astype(str).str.strip().str.upper()
            df["_ORDEN_DOC"] = np.select(
                [
                    doc.str.contains("FACTURA ELECTRÓNICA|FACTURA ELECTRONICA"),
                    doc.str.contains("DOCUMENTO EQUIVALENTE"),
                    doc.str.contains("APPLICATION RESPONSE"),
                ],
                [0, 1, 2], default=3,
            )
            sort_cols.append(col_tipo_doc)
        else:
            df["_ORDEN_DOC"] = 0

        df = df.sort_values(by=sort_cols, kind="stable").reset_index(drop=True)
        return df.drop(columns=["_ORDEN_GRUPO", "_ORDEN_DOC"])

    @staticmethod
    def _cruzar_datos(df_principal, df_conta, df_tercer):
        df_principal = df_principal.drop(
            columns=[c for c in df_principal.columns if str(c).upper() in COLUMNAS_SOBRESCRIBIR], errors='ignore'
        )

        # Emitido → cruza por NIT Receptor; Recibido u otro → NIT Emisor
        es_emitido = df_principal['Grupo'].fillna('').astype(str).str.strip().str.upper() == 'EMITIDO'
        df_principal['NIT_Cruce'] = np.where(
            es_emitido, limpiar_codigo(df_principal['NIT Receptor']), limpiar_codigo(df_principal['NIT Emisor'])
        )

        df_conta = df_conta.assign(NIT_clean=limpiar_codigo(df_conta['NIT'])).drop_duplicates(subset=['NIT_clean'])
        cols_a_traer = ['NIT_clean'] + [c for c in ('TIPO-DETALLE', 'TIPO') if c in df_conta.columns]

        df = df_principal.merge(df_conta[cols_a_traer], left_on='NIT_Cruce', right_on='NIT_clean', how='left')
        for col in ('TIPO-DETALLE', 'TIPO'):
            df[col] = df[col].fillna('') if col in df.columns else ''

        detalle = df['TIPO-DETALLE'].astype(str).str.strip()
        df['CONCEPTO'] = np.where(detalle == '', 'CONCEPTO NO CREADO', detalle)

        nits_terceros = set(limpiar_codigo(df_tercer['NIT']).unique())
        df['TERCERO'] = np.where(df['NIT_Cruce'].isin(nits_terceros) & (df['NIT_Cruce'] != ''), 'CREADO', 'NO CREADO')
        return df

    @staticmethod
    def _calcular_campos(df):
        df['Num.Ext'] = limpiar_codigo(df['Prefijo'], strip=False) + limpiar_codigo(df['Folio'], strip=False)
        df['BASE'] = np.where(df['Total'] == 0, 0, (df['Total'] - df['IVA']))
        return df['TIPO'].astype(str).str.strip().str.upper().str.contains('PERSONAL', na=False, regex=False)

    def _info_tabla_existente(self, nombre_hoja):
        try:
            wb = load_workbook(self.file_path)
            try:
                if nombre_hoja in wb.sheetnames and wb[nombre_hoja].tables:
                    nombre, tbl = next(iter(wb[nombre_hoja].tables.items()))
                    return nombre, tbl.headerRowCount, tbl.totalsRowCount, tbl.tableStyleInfo
            finally:
                wb.close()
        except Exception:
            pass
        return None

    def _escribir_excel(self, nombre_principal, df_export, es_personales):
        try:
            tabla_info = self._info_tabla_existente(nombre_principal)
            n_cols = len(df_export.columns)

            with escritor_excel(self.file_path) as writer:
                df_export.to_excel(writer, index=False, sheet_name=nombre_principal)
                sheet = writer.sheets[nombre_principal]

                personal_fill = PatternFill(start_color="FFCCFF", end_color="FFCCFF", fill_type="solid")
                red_font = Font(color="FF0000")
                col_tipo_doc = buscar_columna(df_export.columns, "TIPO DE DOCUMENTO", "TIPO DOCUMENTO")
                es_nota_credito = (
                    df_export[col_tipo_doc].map(lambda v: "nota de credito" in normalizar_ascii(v) if pd.notna(v) and v else False)
                    if col_tipo_doc else pd.Series(False, index=df_export.index)
                )

                for row_idx, (es_pers, es_nc) in enumerate(zip(es_personales, es_nota_credito), start=2):
                    if not (es_pers or es_nc):
                        continue
                    for col_idx in range(1, n_cols + 1):
                        cell = sheet.cell(row=row_idx, column=col_idx)
                        if es_pers:
                            cell.fill = personal_fill
                        if es_nc:
                            cell.font = red_font

                header_fill = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
                header_font = Font(bold=True, color="FFFFFF")
                for cell in sheet[1]:
                    cell.fill = header_fill
                    cell.font = header_font
                sheet.freeze_panes = 'A2'

                if tabla_info:
                    nombre_tabla, header_rc, totals_rc, style_info = tabla_info
                    tabla = Table(displayName=nombre_tabla, ref=f"A1:{get_column_letter(n_cols)}{len(df_export) + 1}")
                    tabla.headerRowCount = header_rc
                    tabla.totalsRowCount = totals_rc
                    tabla.tableStyleInfo = style_info
                    sheet.add_table(tabla)

                for idx, col_name in enumerate(df_export.columns, start=1):
                    letra = get_column_letter(idx)
                    largo_datos = df_export[col_name].map(lambda v: len(str(v)) if pd.notna(v) and v != '' and v != 0 else 0).max()
                    largo = max(len(str(col_name)), int(largo_datos) if pd.notna(largo_datos) else 0)
                    sheet.column_dimensions[letra].width = max(largo + 2, 10)
                    if col_name in COLUMNAS_OCULTAS:
                        sheet.column_dimensions[letra].hidden = True

        except (ErrorUsuario, ErrorSistema):
            raise
        except Exception as e:
            raise ErrorSistema(f"Error escribiendo la hoja Token: {e}") from e

